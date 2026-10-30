from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


OPERATION_TITLES = {
    "INCOME": "Доход",
    "EXPENSE": "Расход",
    "SAVING": "Накопления",
}

SAVING_TITLES = {
    "DEPOSIT": "Пополнение",
    "WITHDRAW": "Снятие",
}

CATEGORY_TITLES = {
    "INCOME": "Доходы",
    "EXPENSE": "Расходы",
    "SAVING": "Накопления",
}


def style_header(worksheet) -> None:
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(horizontal="center")


def set_column_widths(worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width


def add_excel_table(worksheet, name: str) -> None:
    if worksheet.max_row < 2:
        return

    table = Table(
        displayName=name,
        ref=worksheet.dimensions,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    worksheet.add_table(table)


def build_excel_export(
    transactions: list,
    categories: list,
    currency: str,
) -> BytesIO:
    workbook = Workbook()

    # Лист 1. История транзакций.
    ws_transactions = workbook.active
    ws_transactions.title = "Транзакции"

    ws_transactions.append([
        "Дата и время",
        "Тип операции",
        "Тип накопления",
        "Категория",
        "Сумма",
        "Валюта",
        "Описание",
    ])

    income = Decimal("0")
    expense = Decimal("0")
    deposits = Decimal("0")
    withdrawals = Decimal("0")

    for transaction, category_name in transactions:
        operation_type = transaction.operation_type
        saving_type = transaction.saving_operation_type
        amount = Decimal(transaction.amount)

        operation_date = transaction.operation_date

        if operation_date.tzinfo is not None:
            operation_date = operation_date.replace(tzinfo=None)

        ws_transactions.append([
            operation_date,
            OPERATION_TITLES.get(operation_type, operation_type),
            SAVING_TITLES.get(saving_type, "") if saving_type else "",
            category_name,
            float(amount),
            currency,
            transaction.description or "",
        ])

        if operation_type == "INCOME":
            income += amount
        elif operation_type == "EXPENSE":
            expense += amount
        elif operation_type == "SAVING":
            if saving_type == "DEPOSIT":
                deposits += amount
            elif saving_type == "WITHDRAW":
                withdrawals += amount

    for row in ws_transactions.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY HH:MM:SS"
        row[4].number_format = '#,##0.00'

    ws_transactions.freeze_panes = "A2"
    ws_transactions.auto_filter.ref = ws_transactions.dimensions

    style_header(ws_transactions)
    set_column_widths(
        ws_transactions,
        [20, 18, 18, 25, 16, 12, 40],
    )
    add_excel_table(ws_transactions, "TransactionsTable")

    # Лист 2. Сводка.
    ws_summary = workbook.create_sheet("Сводка")

    net = income - expense
    free_change = income - expense - deposits + withdrawals
    savings_change = deposits - withdrawals

    summary_rows = [
        ["Показатель", "Сумма", "Валюта"],
        ["Доходы", float(income), currency],
        ["Расходы", float(expense), currency],
        ["Пополнения накоплений", float(deposits), currency],
        ["Снятия накоплений", float(withdrawals), currency],
        ["Чистый результат (доходы - расходы)", float(net), currency],
        ["Изменение свободных средств", float(free_change), currency],
        ["Изменение накоплений", float(savings_change), currency],
    ]

    for row in summary_rows:
        ws_summary.append(row)

    for row in ws_summary.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = '#,##0.00'

    style_header(ws_summary)
    set_column_widths(ws_summary, [45, 20, 12])
    ws_summary.freeze_panes = "A2"

    # Лист 3. Категории.
    ws_categories = workbook.create_sheet("Категории")

    ws_categories.append([
        "Категория",
        "Тип",
        "Статус",
    ])

    for category in categories:
        ws_categories.append([
            category.name,
            CATEGORY_TITLES.get(
                category.category_type,
                category.category_type,
            ),
            "Активна" if category.is_active else "Неактивна",
        ])

    style_header(ws_categories)
    set_column_widths(ws_categories, [30, 20, 15])
    ws_categories.freeze_panes = "A2"
    add_excel_table(ws_categories, "CategoriesTable")

    # Возвращаем книгу в памяти.
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output
